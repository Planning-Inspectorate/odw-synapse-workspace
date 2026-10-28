"""Map the manual changes made to the 3rd draft comparison sheet onto the latest comparison sheet.

Manual changes = cells (value or fill colour) that differ between the untouched draft 3 (base)
and the edited draft 3. Cells are matched by (case reference, occurrence, column header), never by
cell address, because the latest sheet has a different layout (header row, extra columns, row order).
Case references can repeat, so the Nth occurrence of a reference in the draft is matched to the Nth
occurrence in the latest sheet.

For every manual change the edited value and/or fill colour overwrites the latest sheet cell.
A change log records whether the latest cell still held the base value (script did not touch it)
or differed (script fix and manual edit both touched it - worth a glance).

Before mapping, the draft's row 2 (notes) and row 4 (duplicate header row) are inserted into the
latest sheet, and the draft's styling for rows 2-4 (fills, fonts, borders, merges, heights) is copied
across. Columns that only exist in the latest sheet are kept and take the style of the draft's last column.

Usage: python map_manual_changes_onto_latest_comparison_sheet.py
Outputs (next to this script):
  S62A_Horizon_vs_Spreadsheet_comparison_with_manual_changes.xlsx
  manual_changes_mapping_log.xlsx
"""

from collections import Counter
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
BASE = HERE / "3rd_draft_comparison_sheet_base.xlsx"
EDITED = HERE / "3rd_draft_comparison_sheet_edited.xlsx"
LATEST = HERE.parents[2] / "outputs" / "S62A_Horizon_vs_Spreadsheet_comparison.xlsx"
OUTPUT = HERE / "S62A_Horizon_vs_Spreadsheet_comparison_with_manual_changes.xlsx"
LOG = HERE / "manual_changes_mapping_log.xlsx"

SHEET = "Template"
DRAFT_HEADER_ROW, DRAFT_FIRST_ROW = 3, 5
# After add_draft_header_rows() the latest sheet has the same layout as the draft.
LATEST_HEADER_ROW, LATEST_FIRST_ROW = 3, 5
KEY_COL = 1


def norm(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.replace("\xa0", " ").strip()
    return v


def fill_sig(cell):
    f = cell.fill
    if f is None or f.fill_type is None:
        return ""
    c = f.fgColor
    if c.type == "rgb":
        return str(c.rgb)
    if c.type == "theme":
        return f"theme{c.theme}/{round(c.tint or 0, 2)}"
    return f"indexed{c.indexed}"


def index_rows(ws, first_row):
    """Map (case reference, occurrence) -> row number."""
    seen = Counter()
    rows = {}
    for r in range(first_row, ws.max_row + 1):
        ref = str(norm(ws.cell(r, KEY_COL).value))
        if not ref:
            continue
        seen[ref] += 1
        rows[(ref, seen[ref])] = r
    return rows


def index_cols(ws, header_row):
    cols = {}
    for c in range(1, ws.max_column + 1):
        h = norm(ws.cell(header_row, c).value)
        if h:
            cols.setdefault(h, c)
    return cols


def copy_style(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def add_draft_header_rows(draft_ws, latest_ws):
    """Insert the draft's notes row (2) and duplicate header row (4) and copy rows 2-4 styling."""
    old_heights = {r: d.height for r, d in latest_ws.row_dimensions.items() if d.height is not None}
    latest_ws.insert_rows(2)
    latest_ws.insert_rows(4)
    for r in sorted(old_heights, reverse=True):
        if r >= 2:
            latest_ws.row_dimensions[r + (1 if r == 2 else 2)].height = old_heights[r]
    for r in (2, 3, 4):
        latest_ws.row_dimensions[r].height = draft_ws.row_dimensions[r].height

    draft_cols = index_cols(draft_ws, DRAFT_HEADER_ROW)
    last_draft_col = draft_ws.max_column
    for c in range(1, latest_ws.max_column + 1):
        header = norm(latest_ws.cell(3, c).value)
        dc = draft_cols.get(header, last_draft_col)
        for r in (2, 3, 4):
            copy_style(draft_ws.cell(r, dc), latest_ws.cell(r, c))
        if header in draft_cols:
            latest_ws.cell(2, c).value = draft_ws.cell(2, dc).value
        latest_ws.cell(4, c).value = latest_ws.cell(3, c).value

    for m in draft_ws.merged_cells.ranges:
        if 2 <= m.min_row <= 4:
            assert draft_cols.get(norm(latest_ws.cell(3, m.min_col).value)) == m.min_col, m
            latest_ws.merge_cells(str(m))

    latest_ws.freeze_panes = "A4"
    latest_ws.auto_filter.ref = f"A3:{get_column_letter(latest_ws.max_column)}{latest_ws.max_row}"


def main():
    base_ws = openpyxl.load_workbook(BASE)[SHEET]
    edit_ws = openpyxl.load_workbook(EDITED)[SHEET]
    latest_wb = openpyxl.load_workbook(LATEST)
    latest_ws = latest_wb[SHEET]
    add_draft_header_rows(edit_ws, latest_ws)

    base_rows, edit_rows = index_rows(base_ws, DRAFT_FIRST_ROW), index_rows(edit_ws, DRAFT_FIRST_ROW)
    latest_rows = index_rows(latest_ws, LATEST_FIRST_ROW)
    base_cols, edit_cols = index_cols(base_ws, DRAFT_HEADER_ROW), index_cols(edit_ws, DRAFT_HEADER_ROW)
    latest_cols = index_cols(latest_ws, LATEST_HEADER_ROW)

    log = []
    for key, er in edit_rows.items():
        br = base_rows.get(key)
        if br is None:
            continue
        for header, ec in edit_cols.items():
            bc = base_cols.get(header)
            if bc is None:
                continue
            e_cell, b_cell = edit_ws.cell(er, ec), base_ws.cell(br, bc)
            value_changed = norm(e_cell.value) != norm(b_cell.value)
            fill_changed = fill_sig(e_cell) != fill_sig(b_cell)
            if not (value_changed or fill_changed):
                continue

            lr, lc = latest_rows.get(key), latest_cols.get(header)
            if lr is None or lc is None:
                if value_changed:
                    log.append((key[0], key[1], header, "Value", b_cell.value, e_cell.value, "", "NOT APPLIED - case/column not in latest sheet"))
                if fill_changed:
                    log.append((key[0], key[1], header, "Fill colour", fill_sig(b_cell) or "none", fill_sig(e_cell) or "none", "", "NOT APPLIED - case/column not in latest sheet"))
                continue
            l_cell = latest_ws.cell(lr, lc)

            if value_changed:
                latest_before = l_cell.value
                same = norm(latest_before) == norm(b_cell.value)
                status = "Applied" if same else "Applied (latest differed from base - script fix also changed this cell)"
                l_cell.value = e_cell.value
                log.append((key[0], key[1], header, "Value", b_cell.value, e_cell.value, latest_before, status))
            if fill_changed:
                latest_before = fill_sig(l_cell) or "none"
                same = latest_before == (fill_sig(b_cell) or "none")
                status = "Applied" if same else "Applied (latest fill differed from base)"
                l_cell.fill = copy(e_cell.fill)
                log.append((key[0], key[1], header, "Fill colour", fill_sig(b_cell) or "none", fill_sig(e_cell) or "none", latest_before, status))

    latest_wb.save(OUTPUT)

    wb = openpyxl.Workbook()
    summary = wb.active
    summary.title = "Summary"
    counts = Counter((r[3], r[7].split(" (")[0].split(" -")[0]) for r in log)
    summary.append(["Change type", "Result", "Count"])
    for (ctype, res), n in sorted(counts.items()):
        summary.append([ctype, res, n])
    overlap = sum(1 for r in log if "differed" in r[7])
    summary.append([])
    summary.append(["Applied where latest differed from base (review)", "", overlap])
    detail = wb.create_sheet("Mapping log")
    detail.append(["Case reference", "Occurrence", "Column", "Change type", "Base", "Manual (edited)", "Latest before merge", "Status"])
    for r in log:
        detail.append(list(r))
    for sheet in (summary, detail):
        for c in sheet[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="305496")
        for i in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(i)].width = 32
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    wb.save(LOG)

    print(f"Wrote {OUTPUT.name} and {LOG.name}")
    for (ctype, res), n in sorted(counts.items()):
        print(f"{ctype}: {res} = {n}")
    print(f"Applied where latest differed from base: {overlap}")


if __name__ == "__main__":
    main()
