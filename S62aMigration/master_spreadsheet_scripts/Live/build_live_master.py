
# %% Imports and paths
import argparse
import csv
import re
from copy import copy
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook


def find_root(start: Path) -> Path:
    """Walk up from the script until the S62aMigration folder (contains csv_and_xlsx_files)."""
    for p in [start, *start.parents]:
        if (p / "csv_and_xlsx_files").is_dir():
            return p
    raise FileNotFoundError("Could not find S62aMigration root (folder containing csv_and_xlsx_files)")


try:
    ROOT = find_root(Path(__file__).resolve().parent)
except (NameError, FileNotFoundError):
    ROOT = Path.cwd()

LIVE_DIR = ROOT / "csv_and_xlsx_files" / "Live"
DEFAULTS = {
    "new_cases": LIVE_DIR / "Master record post June 2026 new cases S62A.xlsx",
    "horizon":   LIVE_DIR / "20261009_query_s62A_extended_tables.csv",
    "template":  ROOT / "outputs" / "MASTER LEGACY cases S62A - final.xlsx",
    "output":    ROOT / "outputs" / "MASTER LIVE cases S62A - with extended data.xlsx",
    "report":    ROOT / "outputs" / "live_master_build_report.csv",
}

NEW_SHEET       = "Template"
TEMPLATE_SHEET  = "Final selected"
HEADER_ROW      = 2          # field names (row 1 holds the group headings)
FIRST_DATA_ROW  = 3
KEEP_TBC_ROWS   = True       # True: keep 'TBC' rows, the transform creates refs for them. False: drop them

COLUMN_RENAMES = {"Pre-application or application": "Application phase"}

# Horizon CaseNodeId per live case, from Sophie (Oct 2026). Joins exactly to CaseNodeId in the extended extract.
# Add new cases here as they come in.
CASE_NODE_IDS = {
    "S62A/2026/0154": "65783528",
    "S62A/2026/0155": "65813466",
    "S62A/2026/0157": "65852926",
    "S62A/2026/0158": "65853590",
    "S62A/2026/0159": "65861567",
    "S62A/2026/0160": "65865261",
    "S62A/2026/0161": "65889926",
    "S62A/2026/0162": "65941707",
    "S62A/2026/0164": "65964609",
    "S62A/2026/0165": "65975066",
    "S62A/2026/0166": "66025179",
    "S62A/2026/0167": "66111347",
    "S62A/2026/0168": "66128256",
    "S62A/2026/0169": "66156691",
    "S62A/2026/0170": "66157244",
    "S62A/2026/0171": "66159640",
}

# Horizon extended columns brought across (names already match the master template)
HORIZON_COLUMNS = [
    # case totals
    "TotalProposedResidentialUnits", "NetResidentialUnitChange", "NetAdditionalFloorspaceSqM",
    "NonResExistingFloorspaceSqM", "NonResFloorspaceLostSqM", "NonResFloorspaceProposedSqM", "NonResNetFloorspaceSqM",
    # waste
    "ActivitiesProcessesDescription", "IsWasteDevelopment", "WasteType", "WasteMaxThroughputType",
    "VoidCapacityCubicMetres", "LiquidWasteMaxAnnualThroughput",
    # parking
    "ParkingGroupID", "ParkingVehicleType", "ParkingExistingSpaces", "ParkingProposedSpaces", "ParkingSpaceDifference",
    # non-residential floorspace
    "NonResGroupID", "UseClass", "FloorspaceChange", "ExistingFloorspaceSqM", "FloorspaceLostSqM", "FloorspaceProposedSqM",
    # housing
    *[f"{h}Housing{x}" for h in ("Existing", "Proposed")
      for x in ("GroupID", "OccupancyType", "UnitType", "1BedCount", "2BedCount", "3BedCount",
                "4PlusBedCount", "UnknownBedCount", "TotalUnits")],
]

REF = re.compile(r"(?i)^\s*S62A\s*/\s*(\d{4})\s*/\s*(\d{7}|\d{4})(RD)?")
report = []


def note(case_ref, area, issue, before=None, after=None):
    report.append({"caseReference": case_ref, "area": area, "issue": issue,
                   "before": None if before is None else str(before), "after": None if after is None else str(after)})


def blank(v):
    return v is None or (isinstance(v, float) and pd.isna(v)) or str(v).strip() in ("", "nan", "None")


def norm_header(v):
    return re.sub(r"\s+", " ", str(v)).strip() if not blank(v) else None


def clean_ref(raw):
    """'S62A/2026/0161/(pre-app)  ON HOLD' -> ('S62A/2026/0161', True). Non-refs (TBC) returned unchanged."""
    if blank(raw):
        return None, False
    text = str(raw)
    on_hold = bool(re.search(r"(?i)on\s*hold", text))
    m = REF.match(text)
    if not m:
        return text.strip(), on_hold
    year, num, rd = m.groups()
    return f"S62A/{year}/{num}{(rd or '').upper()}", on_hold


def as_number(v):
    """CSV text -> int/float so Excel stores numbers, anything else unchanged."""
    if blank(v):
        return None
    s = str(v).strip()
    if re.fullmatch(r"-?\d+", s):
        return int(s)
    if re.fullmatch(r"-?\d*\.\d+", s):
        return float(s)
    return s


# %% 1. New cases sheet
def read_new_cases(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[NEW_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    headers = [norm_header(h) for h in rows[HEADER_ROW - 1]]
    headers = [COLUMN_RENAMES.get(h, h) for h in headers]
    for old, new in COLUMN_RENAMES.items():
        if new in headers:
            note(None, "columns", "Column renamed to match master", old, new)

    dups = sorted({h for h in headers if h and headers.count(h) > 1})
    if dups:
        note(None, "columns", "Duplicate column names in new cases sheet, first non-blank value kept", dups)

    ref_col = headers.index("Case reference")
    cases = []
    for r in rows[FIRST_DATA_ROW - 1:]:
        if all(blank(v) for v in r):
            continue
        if norm_header(r[ref_col]) == "Case reference":
            note(None, "rows", "Example row under headers removed", r[ref_col])
            continue
        rec = {}
        for h, v in zip(headers, r):   # duplicate headers: keep the first non-blank value
            if h and (h not in rec or blank(rec[h])):
                rec[h] = v
        original = rec["Case reference"]
        cleaned, on_hold = clean_ref(original)
        if cleaned != (str(original).strip() if not blank(original) else None):
            note(cleaned, "case reference", "Reference cleaned", original, cleaned)
        if on_hold:
            note(cleaned, "case reference", "Marked ON HOLD in source reference, check status column", original)
        if not REF.match(cleaned or ""):
            if not KEEP_TBC_ROWS:
                note(cleaned, "rows", "No case reference, row dropped (KEEP_TBC_ROWS = False)", original)
                continue
            note(cleaned, "case reference", "No case reference, transform will create one", original)
        rec["Case reference"] = cleaned
        cases.append(rec)
    return cases


# %% 2. Horizon extended data
def read_horizon(path):
    try:
        return pd.read_csv(path, dtype=str, encoding="utf-8")
    except UnicodeDecodeError:
        return pd.read_csv(path, dtype=str, encoding="cp1252")


def map_horizon_ids(hzn, cases):
    """CaseUniqueId -> cleaned case reference, for the live cases only (left join).
    Joined on CaseNodeId using CASE_NODE_IDS."""
    live_refs = {c["Case reference"] for c in cases if REF.match(c["Case reference"] or "")}
    node_to_uid = hzn.groupby("CaseNodeId")["CaseUniqueId"].unique()
    id_to_ref = {}
    for ref in sorted(live_refs):
        node = CASE_NODE_IDS.get(ref)
        if node is None:
            note(ref, "horizon", "No CaseNodeId for this case in CASE_NODE_IDS, no extended data linked")
            continue
        uids = list(node_to_uid.get(node, []))
        if len(uids) != 1:
            note(ref, "horizon", f"CaseNodeId {node} found {len(uids)} times in Horizon extract, not linked", uids)
            continue
        id_to_ref[uids[0]] = ref
        note(ref, "horizon", "Linked to Horizon case on CaseNodeId", node, uids[0])
    for ref in set(CASE_NODE_IDS) - live_refs:
        note(ref, "horizon", "In CASE_NODE_IDS but not in live sheet, ignored")
    return id_to_ref


# Repeating groups in the Horizon extract: one row per group combination. Grouped per case into the
# legacy master format: one value if the case has one group, '[a, b]' (one entry per group, blank kept
# so positions stay aligned) if several. Same format the transform already reads for legacy cases.
GROUP_SETS = {
    "ParkingGroupID": ["ParkingVehicleType", "ParkingExistingSpaces", "ParkingProposedSpaces", "ParkingSpaceDifference"],
    "NonResGroupID": ["UseClass", "FloorspaceChange", "ExistingFloorspaceSqM", "FloorspaceLostSqM", "FloorspaceProposedSqM"],
    **{f"{h}HousingGroupID": [f"{h}Housing{x}" for x in ("OccupancyType", "UnitType", "1BedCount", "2BedCount",
                                                         "3BedCount", "4PlusBedCount", "UnknownBedCount", "TotalUnits")]
       for h in ("Existing", "Proposed")},
    None: ["WasteType", "WasteMaxThroughputType", "VoidCapacityCubicMetres", "LiquidWasteMaxAnnualThroughput"],
}
CASE_LEVEL = ["TotalProposedResidentialUnits", "NetResidentialUnitChange", "NetAdditionalFloorspaceSqM",
              "NonResExistingFloorspaceSqM", "NonResFloorspaceLostSqM", "NonResFloorspaceProposedSqM",
              "NonResNetFloorspaceSqM", "ActivitiesProcessesDescription", "IsWasteDevelopment"]


def _as_list(values):
    """[v] -> v ; [a, b] -> '[a, b]' with blanks left empty so positions line up."""
    if len(values) == 1:
        return values[0]
    return "[" + ", ".join("" if v is None else str(v) for v in values) + "]"


def horizon_by_case(hzn, id_to_ref, template_headers):
    """One dict of extended values per linked case (left join: live cases only)."""
    missing = [c for c in HORIZON_COLUMNS if c not in template_headers]
    if missing:
        note(None, "columns", "Horizon columns not in master template, skipped", missing)
    out = {}
    for cid, g in hzn[hzn["CaseUniqueId"].isin(id_to_ref)].groupby("CaseUniqueId", sort=False):
        ref = id_to_ref[cid]
        rows = [{c: as_number(r.get(c)) for c in HORIZON_COLUMNS} for _, r in g.iterrows()]
        vals = {}
        for c in CASE_LEVEL:   # same on every row: first non-blank
            vals[c] = next((r[c] for r in rows if r[c] is not None), None)
            distinct = {str(r[c]) for r in rows if r[c] is not None}
            if len(distinct) > 1:
                note(ref, "horizon", f"{c} differs across Horizon rows, first kept", sorted(distinct))
        for id_col, fields in GROUP_SETS.items():
            groups = {}
            for r in rows:
                key = r[id_col] if id_col else tuple(r[f] for f in fields)
                if key is None or (not id_col and all(v is None for v in key)):
                    continue
                grp = groups.setdefault(key, {f: None for f in fields})
                for f in fields:
                    if grp[f] is None and r[f] is not None:
                        grp[f] = r[f]
            if not groups:
                continue
            keys = sorted(groups, key=lambda k: (str(type(k)), k))
            if id_col:
                vals[id_col] = _as_list(keys)
            for f in fields:
                col_vals = [groups[k][f] for k in keys]
                if any(isinstance(v, str) and "," in v for v in col_vals):
                    note(ref, "horizon", f"{f} value contains a comma, may split wrongly in the transform", col_vals)
                if any(v is not None for v in col_vals):
                    vals[f] = _as_list(col_vals)
        out[ref] = {k: v for k, v in vals.items() if k in template_headers and v is not None}
        note(ref, "horizon", f"{len(g)} Horizon rows combined onto the case row", after=len(g))
    return out


# %% 3. Write into a copy of the LIVE sheet: same columns, same order, extended columns added on the end
def write_output(new_cases_path, output_path, cases, hzn_by_ref, template_headers):
    wb = load_workbook(new_cases_path)
    ws = wb[NEW_SHEET]
    last_col = max(c for c in range(1, ws.max_column + 1) if norm_header(ws.cell(HEADER_ROW, c).value))
    ref_col = next(c for c in range(1, last_col + 1) if norm_header(ws.cell(HEADER_ROW, c).value) == "Case reference")

    # 1. rename header in place
    for c in range(1, last_col + 1):
        h = norm_header(ws.cell(HEADER_ROW, c).value)
        if h in COLUMN_RENAMES:
            ws.cell(HEADER_ROW, c).value = COLUMN_RENAMES[h]

    # 2. clear the example row (values only, row stays so nothing shifts)
    data_rows = []
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, last_col + 1)]
        if all(blank(v) for v in vals):
            continue
        if norm_header(ws.cell(r, ref_col).value) == "Case reference":
            for c in range(1, last_col + 1):
                ws.cell(r, c).value = None
            continue
        data_rows.append(r)

    # 3. clean refs in place, row by row, in the same order read_new_cases produced them
    kept = [r for r in data_rows]
    refs = [c["Case reference"] for c in cases]
    if not KEEP_TBC_ROWS:
        raise NotImplementedError("KEEP_TBC_ROWS = False is not supported when keeping the live layout")
    assert len(kept) == len(refs), f"row count mismatch {len(kept)} vs {len(refs)}"
    for r, ref in zip(kept, refs):
        ws.cell(r, ref_col).value = ref

    # 4. extended columns added after the last live column, only those the live sheet doesn't already have
    existing = {norm_header(ws.cell(HEADER_ROW, c).value) for c in range(1, last_col + 1)}
    add = [c for c in HORIZON_COLUMNS if c in template_headers and c not in existing]
    head_src = ws.cell(HEADER_ROW, last_col)
    group_src = ws.cell(HEADER_ROW - 1, last_col)
    col_of = {}
    for i, name in enumerate(add, start=last_col + 1):
        h = ws.cell(HEADER_ROW, i, name)
        h.font, h.fill, h.border, h.alignment = copy(head_src.font), copy(head_src.fill), copy(head_src.border), copy(head_src.alignment)
        g = ws.cell(HEADER_ROW - 1, i)
        g.font, g.fill, g.border = copy(group_src.font), copy(group_src.fill), copy(group_src.border)
        col_of[name] = i
    if add:
        ws.cell(HEADER_ROW - 1, last_col + 1).value = "Horizon extended data"
        note(None, "columns", f"{len(add)} Horizon extended columns added after the last live column", add[0], add[-1])

    # 5. extended values written onto each case's own row (one row per case, no rows added)
    added = 0
    for r, ref in zip(kept, refs):
        extra = hzn_by_ref.get(ref)
        if not extra:
            continue
        for name, value in extra.items():
            if name in col_of:
                ws.cell(r, col_of[name], value)
        added += 1

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return added


# %% Run
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    for k, v in DEFAULTS.items():
        ap.add_argument(f"--{k.replace('_', '-')}", type=Path, default=v)
    a = ap.parse_args(argv)

    for k in ("new_cases", "horizon", "template"):
        if not getattr(a, k).is_file():
            raise FileNotFoundError(f"{k}: {getattr(a, k)}")

    cases = read_new_cases(a.new_cases)
    hzn = read_horizon(a.horizon)
    id_to_ref = map_horizon_ids(hzn, cases)

    template_headers = {norm_header(h) for h in
                        next(load_workbook(a.template, read_only=True)[TEMPLATE_SHEET]
                             .iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))}
    hzn_by_ref = horizon_by_case(hzn, id_to_ref, template_headers)
    rows_written = write_output(a.new_cases, a.output, cases, hzn_by_ref, template_headers)

    with open(a.report, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["caseReference", "area", "issue", "before", "after"])
        w.writeheader()
        w.writerows(report)

    linked = sum(1 for n in report if n["issue"].startswith("Linked"))
    print(f"Cases: {len(cases)} | linked to Horizon: {linked} | extended data written to: "
          f"{rows_written} case rows")
    print(f"Output: {a.output}")
    print(f"Report: {a.report}  ({len(report)} entries)")


if __name__ == "__main__":
    main()
