import os
from copy import copy
from pathlib import Path
import re

from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Alignment
from openpyxl.styles.colors import COLOR_INDEX
from openpyxl.utils import get_column_letter


SCRIPT_DIR = Path(os.path.dirname(os.path.abspath(__file__)))

def find_data_root(start_dir, marker="csv_and_xlsx_files", max_up=4):
	d = start_dir
	for _ in range(max_up + 1):
		if (d / marker).is_dir():
			return d
		parent = d.parent
		if parent == d:
			break
		d = parent
	raise FileNotFoundError(
		f"Could not find a '{marker}' folder above {start_dir} "
		f"(searched {max_up + 1} levels up) - check the script's location."
	)

BASE_DIR = find_data_root(SCRIPT_DIR)
COMPARISON_FILE = (
	BASE_DIR
	/ "csv_and_xlsx_files/S62A_Horizon_vs_Spreadsheet_comparison_coloured_sheet.xlsx"
)
MASTER_TEMPLATE_FILE = (
	BASE_DIR / "csv_and_xlsx_files/MASTER LEGACY cases S62A .xlsx"
)

SPREADSHEET_FILE = BASE_DIR / "outputs/S62A_All_Sheets_migrated.xlsx"
OUTPUT_FILE = BASE_DIR / "outputs/MASTER LEGACY cases S62A - final.xlsx"


SHEET_NAME = "Template"
FINAL_SHEET_NAME = "Final selected"
TEMPLATE_HEADER_ROW = 2
TEMPLATE_FIRST_DATA_ROW = 3
KEY_COLUMN = "Case reference"


COMPARISON_HEADER_ROW = None
COMPARISON_COLOUR_ROW = None
COMPARISON_FIRST_DATA_ROW = None


SITE_BLOCK = ["Site address 1", "Site address 2", "Site town or city", "Site county", "Site post code"]
SITE_TRIGGER = "Site address 1"


_REPORTED = set()
THEME_RGB = []          
OVERRIDE_LOG = []       
SITE_SWAP_LOG = []      

EMPTY_FIELD_FILL = PatternFill(fill_type="solid", start_color="D9D9D9", end_color="D9D9D9")
MERGED_FIELD_FILL = PatternFill(fill_type="solid", start_color="FFF2CC", end_color="FFF2CC")


def _is_blank(value):
	return value is None or (isinstance(value, str) and not value.strip())


def _detect_rows(sheet, max_scan=30):
	"""Header = first row whose first cells start with 'Case reference' (notes rows above can be any height).
	Any repeated header rows straight after it are skipped, so data starts on the first real case row."""
	global COMPARISON_HEADER_ROW, COMPARISON_COLOUR_ROW, COMPARISON_FIRST_DATA_ROW

	def is_header(row):
		for column in range(1, min(sheet.max_column, 5) + 1):
			value = sheet.cell(row, column).value
			if isinstance(value, str) and value.strip().casefold().startswith(KEY_COLUMN.casefold()):
				return True
		return False

	for row in range(1, max_scan + 1):
		if is_header(row):
			COMPARISON_HEADER_ROW = COMPARISON_COLOUR_ROW = row
			first_data = row + 1
			while first_data <= sheet.max_row and is_header(first_data):
				print(f"  skipping repeated header row {first_data}")
				first_data += 1
			COMPARISON_FIRST_DATA_ROW = first_data
			return
	raise KeyError(f"Could not find a {KEY_COLUMN!r} header in the first {max_scan} rows")


def _headers(sheet, row):
	"""FIX 2: a blank header straight after 'X (horizon)' is treated as 'X (spreadsheet)'
	(the spreadsheet site address / post code columns have no header in the sheet)."""
	headers = {}
	previous = None
	for column in range(1, sheet.max_column + 1):
		value = sheet.cell(row, column).value
		if value in (None, ""):
			if previous and re.search(r"\(horizon\)$", previous, re.IGNORECASE):
				inferred = re.sub(r"\(horizon\)$", "(spreadsheet)", previous, flags=re.IGNORECASE)
				if inferred not in headers:
					headers[inferred] = column
					if inferred not in _REPORTED:
						_REPORTED.add(inferred)
						print(f"  inferred blank header in column {get_column_letter(column)} as '{inferred}'")
			previous = None
			continue
		header = str(value).strip()
		headers[header] = column
		previous = header
	return headers


def _load_theme(workbook):
	"""Theme colour order as Excel indexes it: lt1, dk1, lt2, dk2, accent1-6, hlink, folHlink."""
	global THEME_RGB
	theme = workbook.loaded_theme
	if isinstance(theme, bytes):
		theme = theme.decode("utf-8", "ignore")
	scheme = re.search(r"<a:clrScheme.*?</a:clrScheme>", theme or "", re.S)
	colours = {}
	if scheme:
		for tag in ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink"):
			m = re.search(rf"<a:{tag}>.*?(?:srgbClr val=\"|lastClr=\")([0-9A-Fa-f]{{6}})", scheme.group(0), re.S)
			colours[tag] = m.group(1) if m else None
	order = ("lt1", "dk1", "lt2", "dk2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink", "folHlink")
	THEME_RGB = [colours.get(tag) for tag in order]


def _apply_tint(components, tint):
	if not tint:
		return components
	if tint > 0:
		return tuple(round(c + (255 - c) * tint) for c in components)
	return tuple(round(c * (1 + tint)) for c in components)


def _rgb(cell):
	fill = cell.fill
	colour = fill.fgColor
	if fill.fill_type != "solid":
		return None
	if colour.type == "rgb" and colour.rgb:
		value = colour.rgb[-6:]
	elif colour.type == "indexed" and colour.indexed is not None:
		value = COLOR_INDEX[colour.indexed][-6:]
	elif colour.type == "theme" and colour.theme is not None and colour.theme < len(THEME_RGB) and THEME_RGB[colour.theme]:
		value = THEME_RGB[colour.theme]
	else:
		return None
	try:
		components = tuple(int(value[index:index + 2], 16) for index in (0, 2, 4))
	except ValueError:
		return None
	return _apply_tint(components, colour.tint if colour.type == "theme" else 0)


def _is_blue(cell):
	"""Dark or light blue cell fill = manual override from the queries spreadsheet."""
	components = _rgb(cell)
	if components is None:
		return False
	red, green, blue = components
	return blue > red + 30 and blue >= green and (max(components) - min(components)) > 25


def _colour_kind(cell):
	"""Classify the decision fill used in the header row."""
	colour = cell.fill.fgColor
	if (
		cell.fill.fill_type == "solid"
		and colour.type == "theme"
		and colour.theme in (0, 1)
		and colour.tint is not None
		and -0.2 <= colour.tint <= -0.1
	):
		return "grey"
	components = _rgb(cell)
	if components is None:
		return "other"
	red, green, blue = components
	brightness = sum(components) / 3
	spread = max(components) - min(components)
	if spread <= 18 and brightness < 245:
		return "grey"
	if green > red and green > blue:
		return "green" if brightness < 180 else "light_green"
	return "other"


def _case_reference_column(sheet):
	headers = _headers(sheet, COMPARISON_HEADER_ROW)
	if KEY_COLUMN in headers:
		return headers[KEY_COLUMN]
	for header, column in headers.items():
		if header.casefold().startswith(f"{KEY_COLUMN.casefold()} ("):
			return column
	raise KeyError(f"Missing {KEY_COLUMN!r} in comparison row {COMPARISON_HEADER_ROW}")


def _case_rows(sheet):
	key_column = _case_reference_column(sheet)
	rows = {}
	for row in range(COMPARISON_FIRST_DATA_ROW, sheet.max_row + 1):
		case_reference = sheet.cell(row, key_column).value
		if not _is_blank(case_reference):
			rows.setdefault(str(case_reference).strip(), []).append(row)
	return rows


def _select_sources(comparison_sheet):
	source_headers = _headers(comparison_sheet, COMPARISON_HEADER_ROW)
	groups = {}
	for header, column in source_headers.items():
		if header.casefold().startswith(KEY_COLUMN.casefold()):
			continue
		match = re.match(r"^(.*) \((horizon|spreadsheet)\)$", header, re.IGNORECASE)
		if not match:
			continue
		field = match.group(1).strip()
		source = match.group(2).casefold()
		groups.setdefault(field, {})[source] = column

	PREFERRED_UNRESOLVED_ORDER = ("spreadsheet", "horizon")

	selections = []
	for field, sources in groups.items():
		kinds = {
			source: _colour_kind(comparison_sheet.cell(COMPARISON_COLOUR_ROW, column))
			for source, column in sources.items()
		}
		green_sources = [source for source in sources if kinds[source] == "green"]

		if green_sources:
			primary_source = green_sources[0]
			fallback_source = "horizon" if primary_source == "spreadsheet" else "spreadsheet"
			# FIX 3: previously the other source was only used if it was light green, so a green
			# Horizon column with a blank cell threw away the spreadsheet value (e.g. Development
			# description). Now it falls back unless that column is explicitly greyed out.
			fallback = (
				sources.get(fallback_source)
				if fallback_source in sources and kinds.get(fallback_source) != "grey"
				else None
			)
			selections.append({
				"header": field,
				"primary": sources[primary_source],
				"fallback": fallback,
				"unresolved": False,
			})
		else:
			ordered = [s for s in PREFERRED_UNRESOLVED_ORDER if s in sources]
			ordered += [s for s in sources if s not in ordered]
			primary_source = ordered[0]
			fallback_source = ordered[1] if len(ordered) > 1 else None
			selections.append({
				"header": field,
				"primary": sources[primary_source],
				"fallback": sources.get(fallback_source) if fallback_source else None,
				"unresolved": True,
			})
	return selections, groups


def _blue_override(value_sheet, style_sheet, rows, columns):
	"""If any cell for this case in the given source columns is blue, that value wins (even if blank).
	Returns (value, fill, column) or None. Columns are checked in priority order."""
	for column in columns:
		if column is None:
			continue
		for row in rows:
			if _is_blue(style_sheet.cell(row, column)):
				return value_sheet.cell(row, column).value, style_sheet.cell(row, column).fill, column
	return None


def _first_value(sheet, rows, column):
	if column is None:
		return None
	for row in rows:
		value = sheet.cell(row, column).value
		if not _is_blank(value):
			return value
	return None


def _copy_template_header(source_sheet, target_sheet, source_column, target_column):
	source = source_sheet.cell(TEMPLATE_HEADER_ROW, source_column)
	target = target_sheet.cell(TEMPLATE_HEADER_ROW, target_column)
	target.value = source.value
	target._style = copy(source._style)
	target.number_format = source.number_format
	target.alignment = copy(source.alignment)
	target.border = copy(source.border)
	target.fill = copy(source.fill)
	target.font = copy(source.font)


def _copy_style(source, target):
	target._style = copy(source._style)
	target.number_format = source.number_format
	target.font = copy(source.font)
	target.fill = copy(source.fill)
	target.border = copy(source.border)
	target.alignment = copy(source.alignment)
	target.protection = copy(source.protection)


def _merged_anchor(sheet, row, column):
	for merged_range in sheet.merged_cells.ranges:
		if (
			merged_range.min_row <= row <= merged_range.max_row
			and merged_range.min_col <= column <= merged_range.max_col
		):
			return merged_range.min_row, merged_range.min_col
	return row, column


def _merge_repeated_row_one(sheet):
	start_column = 1
	while start_column <= sheet.max_column:
		label = sheet.cell(1, start_column).value
		end_column = start_column
		while (
			end_column + 1 <= sheet.max_column
			and sheet.cell(1, end_column + 1).value == label
		):
			end_column += 1
		if label not in (None, "") and end_column > start_column:
			top_left = sheet.cell(1, start_column)
			top_left.alignment = Alignment(horizontal="center", vertical="center")
			sheet.merge_cells(start_row=1, start_column=start_column, end_row=1, end_column=end_column)
		else:
			end_column = start_column
		start_column = end_column + 1


def _load_spreadsheet_site_values():
	"""Site fields per case straight from the migrated spreadsheet output, for spreadsheet
	site columns the comparison sheet doesn't carry (Site address 2, town, county)."""
	if not SPREADSHEET_FILE.is_file():
		print(f"  spreadsheet output not found, site fallback limited to comparison sheet: {SPREADSHEET_FILE}")
		return {}
	wb = load_workbook(SPREADSHEET_FILE, data_only=True)
	try:
		ws = wb[SHEET_NAME]
		headers = {
			str(ws.cell(TEMPLATE_HEADER_ROW, c).value).strip(): c
			for c in range(1, ws.max_column + 1)
			if not _is_blank(ws.cell(TEMPLATE_HEADER_ROW, c).value)
		}
		if KEY_COLUMN not in headers:
			print(f"  {KEY_COLUMN!r} not found in {SPREADSHEET_FILE.name}, site fallback limited to comparison sheet")
			return {}
		missing = [f for f in SITE_BLOCK if f not in headers]
		if missing:
			print(f"  site fields missing from {SPREADSHEET_FILE.name}: {missing}")
		key_column = headers[KEY_COLUMN]
		site_values = {}
		for row in range(TEMPLATE_FIRST_DATA_ROW, ws.max_row + 1):
			case_reference = ws.cell(row, key_column).value
			if _is_blank(case_reference):
				continue
			existing = site_values.setdefault(str(case_reference).strip(), {})
			for field in SITE_BLOCK:
				if field not in headers:
					continue
				value = ws.cell(row, headers[field]).value
				if _is_blank(existing.get(field)) and not _is_blank(value):
					existing[field] = value
		return site_values
	finally:
		wb.close()


def _site_block_values(sheet, source_rows, groups, case_reference, spreadsheet_site):
	"""If Site address 1 (horizon) is blank, return the whole site block from the spreadsheet
	(comparison sheet first, then the migrated spreadsheet output). Blanks are kept so the
	two sources are never mixed. Otherwise return None so site fields go through normal selection."""
	horizon_line_1 = _first_value(sheet, source_rows, groups.get(SITE_TRIGGER, {}).get("horizon"))
	if not _is_blank(horizon_line_1):
		return None
	values = {}
	for field in SITE_BLOCK:
		value = _first_value(sheet, source_rows, groups.get(field, {}).get("spreadsheet"))
		if _is_blank(value):
			value = spreadsheet_site.get(case_reference, {}).get(field)
		values[field] = value
	return values


def extract_final_columns():
	comparison_workbook = load_workbook(COMPARISON_FILE, data_only=True)
	comparison_styles_workbook = load_workbook(COMPARISON_FILE, data_only=False)
	template_workbook = load_workbook(MASTER_TEMPLATE_FILE, data_only=False)
	try:
		comparison_sheet = comparison_workbook[SHEET_NAME]
		comparison_styles_sheet = comparison_styles_workbook[SHEET_NAME]
		template_sheet = template_workbook[SHEET_NAME]
		_load_theme(comparison_styles_workbook)
		_detect_rows(comparison_sheet)
		print(f"Comparison header row: {COMPARISON_HEADER_ROW}, first data row: {COMPARISON_FIRST_DATA_ROW}")
		template_headers = _headers(template_sheet, TEMPLATE_HEADER_ROW)
		selections, groups = _select_sources(comparison_styles_sheet)
		case_rows = _case_rows(comparison_sheet)
		spreadsheet_site = _load_spreadsheet_site_values()

		if FINAL_SHEET_NAME in comparison_styles_workbook.sheetnames:
			del comparison_styles_workbook[FINAL_SHEET_NAME]
		final_sheet = comparison_styles_workbook.create_sheet(FINAL_SHEET_NAME)

		_copy_template_header(template_sheet, final_sheet, template_headers[KEY_COLUMN], 1)
		final_sheet.cell(TEMPLATE_HEADER_ROW, 1).value = KEY_COLUMN
		row_one, key_anchor_column = _merged_anchor(
			comparison_styles_sheet, 1, _case_reference_column(comparison_styles_sheet)
		)
		final_sheet.cell(1, 1).value = comparison_styles_sheet.cell(row_one, key_anchor_column).value
		_copy_style(comparison_styles_sheet.cell(row_one, key_anchor_column), final_sheet.cell(1, 1))
		final_sheet.column_dimensions["A"].width = template_sheet.column_dimensions[
			get_column_letter(template_headers[KEY_COLUMN])
		].width

		output_columns = {}
		for output_column, selection in enumerate(selections, start=2):
			output_columns[selection["header"]] = output_column
			row_one, anchor_column = _merged_anchor(comparison_styles_sheet, 1, selection["primary"])
			final_sheet.cell(1, output_column).value = comparison_styles_sheet.cell(row_one, anchor_column).value
			_copy_style(comparison_styles_sheet.cell(row_one, anchor_column), final_sheet.cell(1, output_column))
			final_sheet.cell(TEMPLATE_HEADER_ROW, output_column).value = selection["header"]
			_copy_style(
				comparison_styles_sheet.cell(COMPARISON_COLOUR_ROW, selection["primary"]),
				final_sheet.cell(TEMPLATE_HEADER_ROW, output_column),
			)
			if selection.get("unresolved"):
				final_sheet.cell(TEMPLATE_HEADER_ROW, output_column).fill = copy(MERGED_FIELD_FILL)
			final_sheet.column_dimensions[get_column_letter(output_column)].width = (
				comparison_styles_sheet.column_dimensions[get_column_letter(selection["primary"])].width
			)
		_merge_repeated_row_one(final_sheet)

		output_row = TEMPLATE_FIRST_DATA_ROW
		# One output row per comparison row, so repeating groups (parking, housing, ...) are not lost.
		# Fields identical across a case's rows are shared by all of them; fields that differ stay on their own row.
		row_jobs = [(case, rows, [row]) for case, rows in case_rows.items() for row in rows]
		for case_reference, case_source_rows, own_rows in row_jobs:
			final_sheet.cell(output_row, 1).value = case_reference
			site_values = _site_block_values(
				comparison_sheet, source_rows, groups, case_reference, spreadsheet_site
			)
			if site_values is not None:
				SITE_SWAP_LOG.append(case_reference)
			for output_column, selection in enumerate(selections, start=2):
				target = final_sheet.cell(output_row, output_column)
				other = [c for c in groups.get(selection["header"], {}).values()
				         if c not in (selection["primary"], selection["fallback"])]
				source_rows = own_rows if _varies(
					comparison_sheet, case_source_rows, [selection["primary"], selection["fallback"], *other]
				) else case_source_rows
				# Blue override (dark or light blue cell): that value is used and keeps its blue fill
				override = _blue_override(
					comparison_sheet, comparison_styles_sheet, source_rows,
					[selection["primary"], selection["fallback"], *other],
				)
				if override:
					value, fill, column = override
					target.value = value
					target.fill = copy(fill)
					source = next((s for s, c in groups.get(selection["header"], {}).items() if c == column), "?")
					OVERRIDE_LOG.append((case_reference, selection["header"], source, value))
					continue
				# Horizon Site address 1 blank: whole site block from the spreadsheet
				if selection["header"] in SITE_BLOCK and site_values is not None:
					target.value = site_values[selection["header"]]
					continue
				primary_value = _first_value(comparison_sheet, source_rows, selection["primary"])
				fallback_value = _first_value(comparison_sheet, source_rows, selection["fallback"])
				value = primary_value if not _is_blank(primary_value) else fallback_value
				final_sheet.cell(output_row, output_column).value = value

				if (
					selection.get("unresolved")
					and not _is_blank(primary_value)
					and not _is_blank(fallback_value)
					and str(primary_value).strip() != str(fallback_value).strip()
				):
					final_sheet.cell(output_row, output_column).fill = copy(MERGED_FIELD_FILL)
			output_row += 1

		for output_column in range(2, len(selections) + 2):
			has_data = any(
				not _is_blank(final_sheet.cell(row, output_column).value)
				for row in range(TEMPLATE_FIRST_DATA_ROW, output_row)
			)
			if not has_data:
				final_sheet.cell(1, output_column).fill = copy(EMPTY_FIELD_FILL)
				# (a column with no data at all can't contain a blue override value, but keep any blue fill)
				final_sheet.cell(TEMPLATE_HEADER_ROW, output_column).fill = copy(EMPTY_FIELD_FILL)
				for row in range(TEMPLATE_FIRST_DATA_ROW, output_row):
					if not _is_blue(final_sheet.cell(row, output_column)):
						final_sheet.cell(row, output_column).fill = copy(EMPTY_FIELD_FILL)

		final_sheet.freeze_panes = f"A{TEMPLATE_FIRST_DATA_ROW}"
		final_sheet.auto_filter.ref = (
			f"A{TEMPLATE_HEADER_ROW}:{get_column_letter(len(selections) + 1)}{output_row - 1}"
		)
		OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
		comparison_styles_workbook.save(OUTPUT_FILE)
		print(f"Wrote {output_row - TEMPLATE_FIRST_DATA_ROW} rows ({len(case_rows)} cases) to {OUTPUT_FILE}")
		print(f"Retained {len(selections)} final columns; comparison Template preserved")
		print(f"Site block taken from spreadsheet (Horizon Site address 1 blank): {len(SITE_SWAP_LOG)} cases")
		by_field = {}
		for _, field, source, _ in OVERRIDE_LOG:
			by_field[(field, source)] = by_field.get((field, source), 0) + 1
		print(f"Blue overrides applied: {len(OVERRIDE_LOG)}")
		for (field, source), n in sorted(by_field.items(), key=lambda x: -x[1]):
			print(f"  {n:>4}  {field} ({source})")
	finally:
		comparison_workbook.close()
		comparison_styles_workbook.close()
		template_workbook.close()


if __name__ == "__main__":
	extract_final_columns()