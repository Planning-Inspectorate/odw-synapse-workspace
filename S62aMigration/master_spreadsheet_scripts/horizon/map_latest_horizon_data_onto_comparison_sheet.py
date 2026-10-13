import csv
import os
import tempfile
from collections import defaultdict
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
MIGRATION_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, "..", ".."))
OUTPUT_DIR = os.path.join(MIGRATION_DIR, "outputs")
HORIZON_FILE = os.path.join(OUTPUT_DIR, "MASTER LEGACY cases S62A - with Horizon data.xlsx")
COMPARISON_FILE = os.path.join(OUTPUT_DIR, "S62A_Horizon_vs_Spreadsheet_comparison_2nd_draft_1.xlsx")
UPDATED_FILE = os.path.join(OUTPUT_DIR, "S62A_Horizon_vs_Spreadsheet_comparison_2nd_draft_1_updated.xlsx")
MAPPING_FILE = os.path.join(OUTPUT_DIR, "horizon_field_mapping.csv")
SHEET_NAME = "Template"
KEY_HEADER = "Case reference"
HORIZON_SUFFIX = " (horizon)"
SOURCE_HEADER_ROW = 2
SOURCE_DATA_START = 3
COMPARISON_HEADER_ROW = 4
COMPARISON_DATA_START = 5


def clean_ref(value):
    return str(value).strip() if value is not None and str(value).strip() else None


def has_value(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def load_extended_data_fields():
    if not os.path.isfile(MAPPING_FILE):
        raise FileNotFoundError(f"Extended-data field mapping not found: {MAPPING_FILE}")

    with open(MAPPING_FILE, newline="", encoding="utf-8-sig") as mapping_stream:
        reader = csv.DictReader(mapping_stream)
        required = {"Field", "Source field"}
        if not required.issubset(reader.fieldnames or []):
            raise ValueError(f"{MAPPING_FILE} must contain columns {sorted(required)}")
        fields = {
            row["Field"].strip().casefold()
            for row in reader
            if row.get("Field", "").strip()
            and row.get("Source field", "").strip().casefold().startswith("extended_data.")
        }
    if not fields:
        raise ValueError(f"No extended_data fields found in {MAPPING_FILE}")
    return fields


def read_horizon_rows(source_ws):
    rows = source_ws.iter_rows(min_row=SOURCE_HEADER_ROW, values_only=True)
    header_values = next(rows, ())
    source_columns = defaultdict(list)
    for idx, value in enumerate(header_values):
        if value is not None:
            source_columns[str(value).strip().casefold()].append(idx)
    key_columns = source_columns.get(KEY_HEADER.casefold(), [])
    if not key_columns:
        raise ValueError(f"Horizon sheet has no {KEY_HEADER!r} column")
    source_key_col = key_columns[0]

    records_by_ref = defaultdict(list)
    source_order = []
    duplicate_header_conflicts = []
    for row_values in rows:
        row_values = tuple(row_values)
        width = len(row_values)

        def get(col):
            return row_values[col] if col < width else None

        case_ref = clean_ref(get(source_key_col))
        if not case_ref:
            continue
        record = {}
        for field, columns in source_columns.items():
            values = [get(col) for col in columns]
            populated = [value for value in values if has_value(value)]
            if len(populated) > 1 and any(value != populated[0] for value in populated[1:]):
                duplicate_header_conflicts.append((case_ref, field, populated))
            record[field] = populated[-1] if populated else None
        records_by_ref[case_ref].append(record)
        source_order.append(case_ref)

    if duplicate_header_conflicts:
        raise ValueError(
            "Conflicting values in duplicate Horizon headers: "
            f"{duplicate_header_conflicts[:10]}"
        )
    return records_by_ref, source_order, set(source_columns)


def update_comparison():
    if not os.path.isfile(HORIZON_FILE):
        raise FileNotFoundError(f"Horizon source workbook not found: {HORIZON_FILE}")

    extended_fields = load_extended_data_fields()
    if not os.path.isfile(COMPARISON_FILE):
        raise FileNotFoundError(f"Comparison workbook not found: {COMPARISON_FILE}")
    baseline_file = COMPARISON_FILE

    source_wb = load_workbook(HORIZON_FILE, data_only=True, read_only=True)
    output_wb = load_workbook(baseline_file, data_only=False)
    try:
        source_ws = source_wb[SHEET_NAME]
        ws = output_wb[SHEET_NAME]
        records_by_ref, source_order, source_fields = read_horizon_rows(source_ws)

        headers = {
            cell.column: str(cell.value).strip()
            for cell in ws[COMPARISON_HEADER_ROW]
            if cell.value is not None
        }
        key_col = next(
            (col for col, name in headers.items() if name.casefold() == KEY_HEADER.casefold()),
            None,
        )
        if key_col is None:
            raise ValueError(f"Comparison sheet has no {KEY_HEADER!r} header")

        horizon_columns = []
        for col, header in headers.items():
            if not header.casefold().endswith(HORIZON_SUFFIX):
                continue
            field = header[:-len(HORIZON_SUFFIX)].strip()
            field_key = field.casefold()
            if field_key not in extended_fields:
                continue
            if field_key not in source_fields:
                raise ValueError(f"Extended Horizon field {field!r} is missing from the source workbook")
            horizon_columns.append((col, field_key))
        if not horizon_columns:
            raise ValueError("No extended-data '(horizon)' columns found in the comparison sheet")

        target_rows_by_ref = defaultdict(list)
        original_last_row = ws.max_row
        for row_num in range(COMPARISON_DATA_START, original_last_row + 1):
            case_ref = clean_ref(ws.cell(row_num, key_col).value)
            if case_ref:
                target_rows_by_ref[case_ref].append(row_num)
        original_row_count = sum(len(rows) for rows in target_rows_by_ref.values())
        original_max_col = ws.max_column
        horizon_col_numbers = {col for col, _ in horizon_columns}

        preserved_template_cells = {
            (row_num, col): ws.cell(row_num, col).value
            for row_num in range(1, original_last_row + 1)
            for col in range(1, original_max_col + 1)
            if not (row_num >= COMPARISON_DATA_START and col in horizon_col_numbers)
        }
        other_sheet_values = {
            sheet.title: {
                (cell.row, cell.column): cell.value
                for row in sheet.iter_rows()
                for cell in row
                if cell.value is not None
            }
            for sheet in output_wb.worksheets
            if sheet.title != SHEET_NAME
        }

        for case_ref, target_rows in target_rows_by_ref.items():
            source_records = records_by_ref.get(case_ref)
            if source_records is None:
                continue
            for occurrence, row_num in enumerate(target_rows):
                record = source_records[occurrence] if occurrence < len(source_records) else None
                for col, field_key in horizon_columns:
                    ws.cell(row_num, col).value = record.get(field_key) if record else None

        extra_records = []
        occurrence_counts = defaultdict(int)
        for case_ref in source_order:
            occurrence = occurrence_counts[case_ref]
            occurrence_counts[case_ref] += 1
            if occurrence >= len(target_rows_by_ref.get(case_ref, [])):
                extra_records.append((case_ref, records_by_ref[case_ref][occurrence]))

        append_start = ws.max_row + 1
        for offset, (case_ref, record) in enumerate(extra_records):
            row_num = append_start + offset
            style_row = target_rows_by_ref.get(case_ref, [COMPARISON_DATA_START])[0]
            for col in range(1, original_max_col + 1):
                source_cell = ws.cell(style_row, col)
                target_cell = ws.cell(row_num, col)
                if source_cell.has_style:
                    target_cell._style = copy(source_cell._style)
                target_cell.value = None
                target_cell.fill = PatternFill(fill_type=None)
            ws.row_dimensions[row_num].height = ws.row_dimensions[style_row].height
            ws.cell(row_num, key_col).value = case_ref
            for col, field_key in horizon_columns:
                ws.cell(row_num, col).value = record.get(field_key)

        new_last_row = append_start + len(extra_records) - 1 if extra_records else original_last_row
        if ws.auto_filter.ref:
            min_col, min_row, max_col, _ = range_boundaries(ws.auto_filter.ref)
            ws.auto_filter.ref = (
                f"{get_column_letter(min_col)}{min_row}:"
                f"{get_column_letter(max_col)}{new_last_row}"
            )

        output_dir = os.path.dirname(UPDATED_FILE)
        fd, candidate_file = tempfile.mkstemp(
            prefix=".S62A_extended_update_", suffix=".xlsx", dir=output_dir
        )
        os.close(fd)
        try:
            output_wb.save(candidate_file)
        finally:
            output_wb.close()
            source_wb.close()

        check_wb = load_workbook(candidate_file, data_only=False)
        try:
            check_ws = check_wb[SHEET_NAME]
            for (row_num, col), value in preserved_template_cells.items():
                if check_ws.cell(row_num, col).value != value:
                    raise AssertionError(f"Non-extended Template cell changed: {get_column_letter(col)}{row_num}")
            for sheet_name, cells in other_sheet_values.items():
                check_sheet = check_wb[sheet_name]
                for (row_num, col), value in cells.items():
                    if check_sheet.cell(row_num, col).value != value:
                        raise AssertionError(f"Other sheet changed: {sheet_name}!{get_column_letter(col)}{row_num}")
            final_rows = sum(
                1 for row_num in range(COMPARISON_DATA_START, check_ws.max_row + 1)
                if clean_ref(check_ws.cell(row_num, key_col).value)
            )
            expected_rows = original_row_count + len(extra_records)
            if final_rows != expected_rows:
                raise AssertionError(f"Unexpected row count: {final_rows}; expected {expected_rows}")
        except Exception:
            check_wb.close()
            os.remove(candidate_file)
            raise
        check_wb.close()

        os.replace(candidate_file, UPDATED_FILE)
        print({
            "source": baseline_file,
            "updated": UPDATED_FILE,
            "extended_fields_updated": len(horizon_columns),
            "source_rows": len(source_order),
            "extra_rows_appended": len(extra_records),
            "comparison_only_cases_preserved": len(set(target_rows_by_ref) - set(records_by_ref)),
        })
    except Exception:
        try:
            output_wb.close()
            source_wb.close()
        except Exception:
            pass
        raise


if __name__ == "__main__":
    update_comparison()